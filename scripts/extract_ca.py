"""Extract continuous-assessment papers (PDF + DOCX) into one corpus text file.

The CAs are the closest thing in the corpus to what students are actually
assessed on. Tutorials show the teaching style; the CAs show the examining
style — phrasing, mark allocation, how much working is expected, which
sub-topics recur year after year. That matters most to the grader, which
otherwise has no idea what "good enough for full marks" looks like here.

Output goes to course_content/ca_papers.txt, which corpus.py loads.

Usage:
    python scripts/extract_ca.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

CONTENT = Path(__file__).resolve().parent.parent / "course_content"
OUT = CONTENT / "ca_papers.txt"


def _pdf(path: Path) -> str:
    import fitz
    with fitz.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _docx(path: Path) -> str:
    from docx import Document
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    # Tables carry MCQ options and mark schemes in several of these papers,
    # and python-docx does not include them in .paragraphs.
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _label(name: str) -> str:
    """A stable, sortable heading: which CA, which sitting."""
    n = name.lower()
    which = "CA2" if re.search(r"ca\s*_?\s*2", n) else "CA1"
    year = (re.search(r"(20\d\d)", n) or [None, "?"])[1]
    month = next((m for m in ("jan", "feb", "mar", "apr", "may", "jun", "jul",
                              "aug", "sep", "oct", "nov", "dec") if m in n), "")
    makeup = " (make-up)" if "makeup" in n.replace("_", "").replace(" ", "") or "make_up" in n or "make up" in n else ""
    mcq = " [MCQ]" if "mcq" in n else ""
    return f"{which} {month.capitalize()} {year}{makeup}{mcq}".replace("  ", " ")


def main() -> None:
    files = sorted(
        [p for p in CONTENT.iterdir()
         if p.suffix.lower() in (".pdf", ".docx")
         and re.match(r"^\d*_?\s*ca\s*_?\d?", p.name, re.I)],
        key=lambda p: (_label(p.name).split()[0], p.name),
    )
    if not files:
        sys.exit("no CA files found in course_content/")

    chunks, skipped = [], []
    for p in files:
        try:
            text = _pdf(p) if p.suffix.lower() == ".pdf" else _docx(p)
        except Exception as e:                       # noqa: BLE001
            skipped.append((p.name, f"{type(e).__name__}: {e}"))
            continue
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        # Two failure modes produce near-empty text and both must be excluded
        # rather than silently padding the corpus: scanned PDFs with no text
        # layer (0 chars), and blank answer *sheets* that carry only headers
        # and "Question 1" placeholders (~400 chars). A genuine paper is
        # several thousand characters, so the threshold separates them cleanly.
        if len(text) < 1500:
            why = ("no text layer — scanned images, needs OCR" if len(text) < 200
                   else f"only {len(text)} chars — blank answer sheet, no questions")
            skipped.append((p.name, why))
            continue
        chunks.append(f"\n{'=' * 72}\n{_label(p.name)}  (source: {p.name})\n{'=' * 72}\n{text}")
        print(f"  {_label(p.name):28} {len(text):>7,} chars   {p.name}")

    OUT.write_text("\n".join(chunks), encoding="utf-8")
    total = sum(len(c) for c in chunks)
    print(f"\n{len(chunks)} papers -> {OUT.name}  ({total:,} chars)")
    if skipped:
        print(f"\nSKIPPED {len(skipped)} — these need OCR or manual conversion:")
        for name, why in skipped:
            print(f"  {name}: {why}")


if __name__ == "__main__":
    main()
